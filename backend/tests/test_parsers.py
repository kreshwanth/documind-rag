import os
import tempfile
import fitz  # PyMuPDF
import docx
from app.services.parser_service import document_parser, PDFParser, DOCXParser, TXTParser

def test_txt_parser():
    with tempfile.NamedTemporaryFile("w", delete=False, suffix=".txt", encoding="utf-8") as f:
        f.write("DocuMind TXT document content. Section 1.\n\nSection 2: Architecture.")
        temp_path = f.name

    try:
        results = document_parser.extract_text(temp_path, "txt")
        assert len(results) == 1
        assert "DocuMind TXT document content" in results[0]["text"]
        assert results[0]["page_number"] is None
    finally:
        os.remove(temp_path)

def test_pdf_parser_page_preservation():
    # Generate a dynamic multi-page test PDF in memory
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "This is content on Page 1 of the enterprise report.")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "This is financial summary on Page 2.")
    
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as f:
        temp_path = f.name
    doc.save(temp_path)
    doc.close()

    try:
        results = document_parser.extract_text(temp_path, "pdf")
        assert len(results) == 2
        assert results[0]["page_number"] == 1
        assert "Page 1" in results[0]["text"]
        assert results[1]["page_number"] == 2
        assert "Page 2" in results[1]["text"]
    finally:
        os.remove(temp_path)

def test_docx_parser():
    doc = docx.Document()
    doc.add_heading("DocuMind Enterprise Specification", 0)
    doc.add_paragraph("Paragraph 1: Core retrieval capabilities.")
    
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as f:
        temp_path = f.name
    doc.save(temp_path)

    try:
        results = document_parser.extract_text(temp_path, "docx")
        assert len(results) >= 1
        assert "Core retrieval capabilities" in results[0]["text"]
    finally:
        os.remove(temp_path)
