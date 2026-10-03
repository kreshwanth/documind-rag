import os
import re
from typing import List, Dict, Any, Optional
import fitz  # PyMuPDF
import docx
import logging

logger = logging.getLogger("documind.parser")

class BaseParser:
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        raise NotImplementedError

class PDFParser(BaseParser):
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        """Extract text from PDF per page preserving exact page numbers and normalizing ligatures."""
        import unicodedata
        pages_content = []
        try:
            doc = fitz.open(file_path)
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                raw_text = page.get_text("text").strip()
                if raw_text:
                    # Normalize ligatures and unicode characters
                    norm_text = unicodedata.normalize("NFKD", raw_text)
                    norm_text = (
                        norm_text.replace("\ufb00", "ff")
                        .replace("\ufb01", "fi")
                        .replace("\ufb02", "fl")
                        .replace("\ufb03", "ffi")
                        .replace("\ufb04", "ffl")
                        .replace("’", "'")
                        .replace("“", '"')
                        .replace("”", '"')
                        .replace("–", "-")
                        .replace("—", "-")
                    )
                    pages_content.append({
                        "page_number": page_idx + 1,
                        "text": norm_text
                    })
            doc.close()
        except Exception as e:
            logger.error(f"PDF extraction error in {file_path}: {e}")
            raise RuntimeError(f"PDF extraction error: {str(e)}")
        return pages_content

class DOCXParser(BaseParser):
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        """Extract text from DOCX documents with page_number=None (or section numbers)."""
        paragraphs_text = []
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                clean = para.text.strip()
                if clean:
                    paragraphs_text.append(clean)
            
            # Extract table contents
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        paragraphs_text.append(" | ".join(row_data))
                        
            full_text = "\n\n".join(paragraphs_text).strip()
            if not full_text:
                return []
            return [{"page_number": None, "text": full_text}]
        except Exception as e:
            logger.error(f"DOCX extraction error in {file_path}: {e}")
            raise RuntimeError(f"DOCX extraction error: {str(e)}")

class TXTParser(BaseParser):
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        """Extract text from plain text files with fallback encodings."""
        content = None
        for enc in ["utf-8", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if content is None:
            raise RuntimeError(f"Could not decode text file {file_path} with standard encodings.")

        clean = content.strip()
        if not clean:
            return []
        return [{"page_number": None, "text": clean}]

class DocumentParser:
    _parsers = {
        "pdf": PDFParser(),
        "docx": DOCXParser(),
        "doc": DOCXParser(),
        "txt": TXTParser(),
        "text": TXTParser(),
        "md": TXTParser()
    }

    @classmethod
    def extract_text(cls, file_path: str, file_type: str) -> List[Dict[str, Any]]:
        file_ext = file_type.lower().strip(".")
        parser = cls._parsers.get(file_ext)
        if not parser:
            raise ValueError(f"Unsupported file format: {file_ext}. Supported: PDF, DOCX, TXT")
        return parser.parse(file_path)

document_parser = DocumentParser()
